"""DocxParser — native Microsoft Word (.docx) document parser using python-docx.

Extracts structured headings, paragraphs, and tables (with HTML representations)
conforming to the RAG ingestion pipeline contract.
"""
from __future__ import annotations

import html
import time
from pathlib import Path

from packages.ingestion.parsed_document import (
    NormalizedDocument,
    ParsedHeading,
    ParsedPage,
    ParsedParagraph,
    ParsedTable,
    ParserProvenance,
)
from packages.ingestion.parsers.base import BaseParser

_DOCX_PARSER_VERSION = "1.0.0"

try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


class DocxParser(BaseParser):
    """Parses .docx files into structured NormalizedDocument representations."""

    @property
    def name(self) -> str:
        return "python_docx"

    @property
    def version(self) -> str:
        return _DOCX_PARSER_VERSION

    def parse(
        self,
        file_path: Path,
        document_id: str,
        document_version_id: str,
    ) -> NormalizedDocument:
        """Parse a Word .docx file into a NormalizedDocument."""
        self._check_path(file_path)

        if not _DOCX_AVAILABLE:
            raise RuntimeError("python-docx is required for DOCX parsing.")

        start = time.monotonic()
        doc = docx.Document(str(file_path))

        headings: list[ParsedHeading] = []
        paragraphs: list[ParsedParagraph] = []
        tables: list[ParsedTable] = []
        reading_order = 0
        current_section = ""

        # Process document body paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = (p.style.name if p.style else "").lower()
            if "heading 1" in style_name or "title" in style_name:
                level = 1
                headings.append(ParsedHeading(level=level, text=text, page=1, path=text))
                current_section = text
            elif "heading 2" in style_name:
                level = 2
                headings.append(ParsedHeading(level=level, text=text, page=1, path=f"{current_section} > {text}" if current_section else text))
                current_section = text
            elif "heading" in style_name:
                level = 3
                headings.append(ParsedHeading(level=level, text=text, page=1, path=f"{current_section} > {text}" if current_section else text))
            else:
                reading_order += 1
                paragraphs.append(
                    ParsedParagraph(
                        text=text,
                        page=1,
                        section_path=current_section,
                        reading_order=reading_order,
                    )
                )

        # Process tables
        for t_idx, t in enumerate(doc.tables, 1):
            reading_order += 1
            rows_data: list[list[str]] = []
            html_rows: list[str] = []
            
            for r_idx, row in enumerate(t.rows):
                row_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_cells)
                
                tag = "th" if r_idx == 0 else "td"
                cells_html = "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row_cells)
                html_rows.append(f"<tr>{cells_html}</tr>")

            plain_text_table = "\n".join(" | ".join(r) for r in rows_data)
            table_html = f"<table>\n" + "\n".join(html_rows) + "\n</table>"
            
            row_count = len(rows_data)
            col_count = len(rows_data[0]) if rows_data else 0

            tables.append(
                ParsedTable(
                    text=plain_text_table,
                    table_html=table_html,
                    caption=f"Table {t_idx}",
                    page=1,
                    section_path=current_section,
                    reading_order=reading_order,
                    row_count=row_count,
                    col_count=col_count,
                )
            )

        duration = time.monotonic() - start

        # Synthesize page
        pages = [
            ParsedPage(
                page_number=1,
                text="\n\n".join(p.text for p in paragraphs),
                is_empty=len(paragraphs) == 0 and len(tables) == 0,
            )
        ]

        return NormalizedDocument(
            document_id=document_id,
            document_version_id=document_version_id,
            pages=pages,
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            provenance=ParserProvenance(
                parser_name=self.name,
                parser_version=self.version,
                duration_seconds=round(duration, 3),
                fallback_used=False,
                fallback_reason="",
            ),
        )
